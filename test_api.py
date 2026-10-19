import os, tempfile
os.environ.setdefault("DATABASE_URL", f"sqlite:///{tempfile.mkdtemp()}/t.db")
os.environ.pop("APP_USER", None); os.environ.pop("APP_PASSWORD", None)
from fastapi.testclient import TestClient
from app import app

c = TestClient(app)

def tree():
    assert c.put("/api/projects/p1", json={"name": "Plant A", "client": "X"}).status_code == 200
    assert c.put("/api/packages/k1", json={"projectId": "p1", "name": "Earthing", "budget": 1000}).status_code == 200
    assert c.put("/api/suppliers/s1", json={"packageId": "k1", "name": "A", "tech": "Compliant", "stage": "Submitted", "total": 900}).status_code == 200

def test_hierarchy_and_history():
    tree()
    c.put("/api/suppliers/s1", json={"packageId": "k1", "name": "A", "tech": "Compliant", "stage": "Awarded", "total": 900})
    s = [x for x in c.get("/api/all").json()["suppliers"] if x["id"] == "s1"][0]
    assert [h["s"] for h in s["history"]] == ["Compliant / Submitted", "Compliant / Awarded"]

def test_foreign_keys_and_validation():
    tree()
    assert c.put("/api/packages/kx", json={"projectId": "nope", "name": "x"}).status_code == 422
    assert c.put("/api/suppliers/sx", json={"packageId": "nope", "name": "x"}).status_code == 422
    assert c.put("/api/suppliers/sy", json={"packageId": "k1", "name": "x", "tech": "Bad"}).status_code == 422
    assert c.put("/api/projects/py", json={"name": ""}).status_code == 422

def test_cascade_delete_and_idempotent():
    tree()
    assert c.delete("/api/projects/p1").status_code == 200
    d = c.get("/api/all").json()
    assert not [x for x in d["packages"] if x["id"] == "k1"] and not [x for x in d["suppliers"] if x["id"] == "s1"]
    assert c.delete("/api/projects/p1").status_code == 200       # second delete is harmless
    assert c.delete("/api/nothing/1").status_code == 404

def test_backup_restore_roundtrip():
    tree()
    b = c.get("/api/backup")
    assert "attachment" in b.headers["content-disposition"]
    data = b.json()
    c.delete("/api/projects/p1")
    assert c.post("/api/restore", json=data).status_code == 200
    s = [x for x in c.get("/api/all").json()["suppliers"] if x["id"] == "s1"][0]
    assert s["total"] == 900 and s["history"]
    assert c.post("/api/restore", json={"packages": [{"id": "z", "projectId": "ghost", "name": "n"}]}).status_code == 422

def test_basic_auth(monkeypatch):
    monkeypatch.setenv("APP_USER", "u"); monkeypatch.setenv("APP_PASSWORD", "p")
    assert c.get("/api/all").status_code == 401
    assert c.get("/api/all", auth=("u", "p")).status_code == 200


def test_notes_crud_scope_and_cascade():
    tree()
    assert c.put("/api/notes/n1", json={"packageId": "k1", "supplierId": "s1", "body": "Called A"}).status_code == 200
    assert c.put("/api/notes/n2", json={"packageId": "k1", "body": "Package-level"}).status_code == 200
    n = {x["id"]: x for x in c.get("/api/all").json()["notes"]}
    assert n["n1"]["supplierId"] == "s1" and n["n2"]["supplierId"] is None and n["n1"]["createdAt"]
    created = n["n1"]["createdAt"]
    c.put("/api/notes/n1", json={"packageId": "k1", "supplierId": "s1", "body": "Edited"})
    assert [x for x in c.get("/api/all").json()["notes"] if x["id"] == "n1"][0]["createdAt"] == created   # edit keeps date
    assert c.put("/api/notes/n3", json={"packageId": "nope", "body": "x"}).status_code == 422
    assert c.put("/api/notes/n4", json={"packageId": "k1", "body": ""}).status_code == 422
    c.put("/api/packages/k2", json={"projectId": "p1", "name": "Other"})
    assert c.put("/api/notes/n5", json={"packageId": "k2", "supplierId": "s1", "body": "wrong pkg"}).status_code == 422
    c.delete("/api/suppliers/s1")
    ids = [x["id"] for x in c.get("/api/all").json()["notes"]]
    assert "n1" not in ids and "n2" in ids            # supplier note cascades, package note stays
    c.delete("/api/projects/p1")
    assert not c.get("/api/all").json()["notes"]

def test_notes_in_backup_restore():
    tree()
    c.put("/api/notes/n1", json={"packageId": "k1", "supplierId": "s1", "body": "Keep me", "createdAt": "2026-10-01T08:00:00"})
    data = c.get("/api/backup").json()
    assert data["notes"]
    c.delete("/api/projects/p1")
    assert c.post("/api/restore", json=data).status_code == 200
    n = [x for x in c.get("/api/all").json()["notes"] if x["id"] == "n1"][0]
    assert n["body"] == "Keep me" and n["createdAt"].startswith("2026-10-01T08:00")

def test_excel_exports():
    from io import BytesIO
    from openpyxl import load_workbook
    tree()
    c.put("/api/suppliers/s2", json={"packageId": "k1", "name": "B", "tech": "Non-compliant", "stage": "Submitted", "total": 500})
    c.put("/api/notes/n1", json={"packageId": "k1", "supplierId": "s1", "body": "Called A"})
    r = c.get("/api/export/package/k1.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"] and ".xlsx" in r.headers["content-disposition"]
    ws = load_workbook(BytesIO(r.content)).active
    assert ws["B4"].value == 1000 and ws["B9"].value == "B" and ws["C9"].value == "A"       # sorted by bid, cheapest first
    assert str(ws["B6"].value).startswith("=") and str(ws["B13"].value).startswith("=")     # formulas, not pasted numbers
    assert any(cell.value == "Called A" for row in ws.iter_rows() for cell in row)
    wb = load_workbook(BytesIO(c.get("/api/export/project/p1.xlsx").content))
    assert wb.sheetnames[0] == "Summary" and "Earthing" in wb.sheetnames
    assert c.get("/api/export/package/none.xlsx").status_code == 404
    assert c.get("/api/export/project/none.xlsx").status_code == 404
